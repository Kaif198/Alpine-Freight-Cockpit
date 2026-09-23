from io import BytesIO
from datetime import datetime
from urllib.parse import urljoin
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from scraper.network import get
from scraper.schemas import Diesel

NAME = 'European Commission Weekly Oil Bulletin'
PAGE = 'https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en'
URL = 'https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx'
DAYS = 7


def fetch() -> bytes:
    # Discover the current history link; do not assume the download UUID is permanent.
    global URL
    soup = BeautifulSoup(get(PAGE), 'html.parser')
    links = [urljoin(PAGE, a['href']) for a in soup.find_all('a', href=True) if 'Weekly_Oil_Bulletin_Prices_History' in a['href']]
    if len(set(links)) != 1:
        raise ValueError('Could not unambiguously identify the official historical workbook')
    URL = links[0]
    return get(URL)


def parse(raw: bytes) -> list[Diesel]:
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook['Prices with taxes']
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    countries = {'AT', 'DE', 'IT', 'PL', 'HU', 'FR', 'NL', 'CZ'}
    columns = [(i, h.split('_')[0]) for i, h in enumerate(headers) if isinstance(h, str) and h.endswith('_price_with_tax_diesel') and h.split('_')[0] in countries]
    if len(columns) != len(countries):
        raise ValueError('Diesel workbook headers changed')
    records = []
    dates = set()
    for row in rows:
        if not isinstance(row[0], datetime):
            continue
        if len(dates) >= 27:
            break
        dates.add(row[0].date())
        for index, country in columns:
            value = row[index]
            if value is None:
                continue
            if isinstance(value, bool) or not isinstance(value, (float, int)):
                raise ValueError(f'Non-numeric diesel value for {country}')
            records.append(Diesel(country=country, date=row[0].date().isoformat(), eur_per_litre=value / 1000))
    workbook.close()
    if not records:
        raise ValueError('No valid diesel observations')
    return sorted(records, key=lambda x: (x.date, x.country))

